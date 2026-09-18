from datetime import datetime,timedelta
from io import BytesIO
from flask import Blueprint,g,request,send_file
from sqlalchemy import text
from .api_common import ApiError,date_value,endpoint,engine,page,response
from .disponibilidad_service import ZONA_NEGOCIO
from .permisos import requiere_roles

reportes=Blueprint('reportes',__name__,url_prefix='/api/v1')


def periodo():
    if set(request.args)-{'desde','hasta'} or any(len(request.args.getlist(k))!=1 for k in request.args):raise ApiError('Usa únicamente desde y hasta sin repetirlos.')
    hoy=datetime.now(ZONA_NEGOCIO).date()
    start=date_value(request.args.get('desde',hoy.replace(day=1).isoformat()));end=date_value(request.args.get('hasta',hoy.isoformat()))
    if end<start or (end-start).days>366:raise ApiError('El período debe ser ordenado y de máximo 366 días.')
    return {'desde':start,'hasta':end,'fin_exclusivo':end+timedelta(days=1)}


@reportes.get('/admin/reportes/resumen')
@requiere_roles('administrador')
@endpoint
def resumen():
    p=periodo()
    with engine().connect() as c:
        estados=[dict(r) for r in c.execute(text('SELECT estado,COUNT(*) AS cantidad FROM citas WHERE fecha BETWEEN :desde AND :hasta GROUP BY estado'),p).mappings()]
        servicios=[dict(r) for r in c.execute(text('SELECT s.id,s.nombre,COUNT(*) AS citas FROM citas c JOIN servicios s ON s.id=c.servicio_id WHERE c.fecha BETWEEN :desde AND :hasta GROUP BY s.id,s.nombre ORDER BY citas DESC,s.id LIMIT 10'),p).mappings()]
        ventas=c.execute(text("SELECT COUNT(*) AS pedidos,COALESCE(SUM(total),0) AS total_demostracion FROM pedidos WHERE fecha>=:desde AND fecha<:fin_exclusivo AND estado='completado'"),p).mappings().one()
        clientes=c.execute(text("SELECT COUNT(*) FROM usuarios WHERE rol='cliente' AND activo=1")).scalar_one()
        inventario=c.execute(text('SELECT COUNT(*) AS articulos,COALESCE(SUM(stock),0) AS unidades FROM productos WHERE activo=1')).mappings().one()
    return response({'periodo':{'desde':p['desde'],'hasta':p['hasta']},'citas_por_estado':estados,'servicios_mas_solicitados':servicios,'ventas_simuladas':dict(ventas),'clientes_activos_actuales':clientes,'inventario_actual':dict(inventario)})


@reportes.get('/admin/reportes/citas.xlsx')
@requiere_roles('administrador')
@endpoint
def exportar_citas():
    p=periodo()
    with engine().connect() as c:
        rows=c.execute(text("""SELECT c.id,u.nombre AS cliente,s.nombre AS servicio,prof.nombre AS profesional,
            DATE_FORMAT(c.fecha,'%Y-%m-%d') AS fecha,TIME_FORMAT(c.hora,'%H:%i') AS hora,c.duracion_min,c.estado
            FROM citas c JOIN usuarios u ON u.id=c.cliente_id JOIN servicios s ON s.id=c.servicio_id
            LEFT JOIN usuarios prof ON prof.id=c.personal_id WHERE c.fecha BETWEEN :desde AND :hasta
            ORDER BY c.fecha,c.hora,c.id LIMIT 10001"""),p).mappings().all()
    if len(rows)>10000:raise ApiError('El reporte excede 10000 citas. Reduce el período.',413)
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();ws=wb.active;ws.title='Citas'
    fields=['id','cliente','servicio','profesional','fecha','hora','duracion_min','estado']
    ws.append(fields)
    for row in rows:
        ws.append([row[k] for k in fields])
        # Los nombres ingresados por usuarios siempre son texto, nunca fórmulas.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='59384F')
    for col,width in zip('ABCDEFGH',[12,32,32,32,16,12,18,18]):ws.column_dimensions[col].width=width
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    output=BytesIO();wb.save(output);output.seek(0)
    return send_file(output,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name='citas.xlsx')


@reportes.get('/recordatorios')
@requiere_roles('cliente','personal')
@endpoint
def recordatorios():
    now=datetime.now(ZONA_NEGOCIO).replace(tzinfo=None)
    field='cliente_id' if g.usuario_actual['rol']=='cliente' else 'personal_id'
    with engine().connect() as c:
        rows=c.execute(text("SELECT id,fecha,TIME_FORMAT(hora,'%H:%i') AS hora,estado FROM citas WHERE "+field+"=:id AND estado IN ('pendiente','confirmada') AND TIMESTAMP(fecha,hora)>:ahora AND TIMESTAMP(fecha,hora)<=:fin ORDER BY fecha,hora LIMIT 100"),{'id':g.usuario_actual['id'],'ahora':now,'fin':now+timedelta(hours=24)}).mappings().all()
    return response({'canal':'en_aplicacion','proximas_24_horas':[dict(r) for r in rows]},message='Recordatorios para consultar dentro de la aplicación')


@reportes.get('/admin/servicios')
@requiere_roles('administrador')
@endpoint
def catalogo_admin():
    with engine().connect() as c:return page(c,'SELECT * FROM servicios ORDER BY id')
